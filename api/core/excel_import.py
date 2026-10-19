"""Dealer Excel import: one row = one car.

`import_cars(org, file, commit=False)` checks every row and reports, per row,
whether it would create a car, update one, or fails (and why). With
`commit=True` it saves the valid rows; rows with errors are skipped.
"""

import re
from dataclasses import dataclass, field
from io import BytesIO

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Car

# Header in the file -> Car field. Headers are matched case-insensitively.
COLUMNS = {
    "ID": "external_id",
    "მარკა": "make",
    "მოდელი": "model",
    "წელი": "year",
    "ძარის ტიპი": "body",
    "საწვავი": "fuel",
    "ძრავა": "engine",
    "გადაცემათა კოლოფი": "transmission",
    "გარბენი": "mileage_km",
    "ფასი": "price_gel",
    "ქალაქი": "city",
}
REQUIRED = {"external_id", "make", "model", "year", "fuel", "transmission", "mileage_km", "price_gel", "city"}

MISSING_MESSAGES = {
    "external_id": "ID აკლია",
    "make": "მარკა აკლია",
    "model": "მოდელი აკლია",
    "year": "წელი აკლია",
    "fuel": "საწვავი აკლია",
    "transmission": "გადაცემათა კოლოფი აკლია",
    "mileage_km": "გარბენი აკლია",
    "price_gel": "ფასი აკლია",
    "city": "ქალაქი აკლია",
}


class ImportFileError(Exception):
    """The file as a whole can't be read (wrong format, missing columns...)."""


@dataclass
class RowResult:
    row: int
    external_id: str
    action: str  # "new" | "updated" | "error"
    errors: list = field(default_factory=list)
    car: str = ""


@dataclass
class ImportResult:
    rows: list
    committed: bool

    @property
    def counts(self):
        out = {"new": 0, "updated": 0, "error": 0}
        for r in self.rows:
            out[r.action] += 1
        return out


def _choice_lookup(choices):
    lookup = {}
    for value, label in choices:
        lookup[value.lower()] = value
        lookup[str(label).lower()] = value
    return lookup


FUEL = _choice_lookup(Car.Fuel.choices)
TRANSMISSION = _choice_lookup(Car.Transmission.choices)
BODY = _choice_lookup(Car.Body.choices)
# "ჯიპი / SUV" is awkward to type; accept the parts too.
BODY.update({"ჯიპი": "suv", "suv": "suv"})


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value):
    """Accepts 85000, 85000.0, "85 000", "85,000 ლ"."""
    if isinstance(value, (int, float)):
        if value < 0 or (isinstance(value, float) and not value.is_integer()):
            return None
        return int(value)
    digits = re.sub(r"[\s,]|ლ|GEL|კმ", "", _text(value), flags=re.IGNORECASE)
    return int(digits) if digits.isdigit() else None


def _parse_row(raw):
    """raw: {field: cell value}. Returns (clean values, errors)."""
    values, errors = {}, []
    for f in REQUIRED:
        if _text(raw.get(f)) == "":
            errors.append(MISSING_MESSAGES[f])

    values["external_id"] = _text(raw.get("external_id"))[:60]
    values["make"] = _text(raw.get("make"))[:60]
    values["model"] = _text(raw.get("model"))[:80]
    values["engine"] = _text(raw.get("engine"))[:40]
    values["city"] = _text(raw.get("city"))[:60]

    if _text(raw.get("year")):
        year = _number(raw.get("year"))
        if year is None or not 1950 <= year <= timezone.now().year + 1:
            errors.append("წელი არასწორია")
        values["year"] = year

    for f, label in (("mileage_km", "გარბენი"), ("price_gel", "ფასი")):
        if _text(raw.get(f)):
            n = _number(raw.get(f))
            if n is None:
                errors.append(f"{label} უნდა იყოს რიცხვი")
            values[f] = n
    if values.get("price_gel") == 0:
        errors.append("ფასი 0-ზე მეტი უნდა იყოს")

    for f, lookup, label in (
        ("fuel", FUEL, "საწვავი"),
        ("transmission", TRANSMISSION, "გადაცემათა კოლოფი"),
    ):
        text = _text(raw.get(f)).lower()
        if text:
            if text not in lookup:
                errors.append(f"{label} „{_text(raw.get(f))}“ უცნობია")
            values[f] = lookup.get(text)

    # Optional column: blank is fine, unknown text is an error.
    body = _text(raw.get("body")).lower()
    if body and body not in BODY:
        errors.append(f"ძარის ტიპი „{_text(raw.get('body'))}“ უცნობია")
    values["body"] = BODY.get(body, "")

    return values, errors


def read_rows(file):
    """Yields (excel row number, {field: value}) for non-empty rows."""
    if getattr(file, "size", 0) > settings.EXCEL_MAX_BYTES:
        raise ImportFileError("ფაილი ძალიან დიდია (მაქს. 5 MB).")
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("ფაილი ვერ წავიკითხეთ. ატვირთეთ .xlsx ფორმატის ფაილი.") from exc
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    by_name = {k.lower(): v for k, v in COLUMNS.items()}
    index = {}
    for i, h in enumerate(header):
        f = by_name.get(_text(h).lower())
        if f and f not in index:
            index[f] = i
    missing = [h for h, f in COLUMNS.items() if f in REQUIRED and f not in index]
    if missing:
        raise ImportFileError("ფაილს აკლია სვეტები: " + ", ".join(missing))

    count = 0
    for n, row in enumerate(rows, start=2):
        if not row or all(_text(c) == "" for c in row):
            continue
        count += 1
        if count > settings.EXCEL_MAX_ROWS:
            raise ImportFileError(f"ერთ ფაილში მაქსიმუმ {settings.EXCEL_MAX_ROWS} მანქანაა დასაშვები.")
        yield n, {f: (row[i] if i < len(row) else None) for f, i in index.items()}


def import_cars(organization, file, commit=False):
    parsed = list(read_rows(file))
    existing = {
        c.external_id: c for c in Car.objects.filter(organization=organization)
    }
    seen, results, to_save = set(), [], []

    for n, raw in parsed:
        values, errors = _parse_row(raw)
        ext = values["external_id"]
        if ext and ext in seen:
            errors.append(f"ID {ext} ფაილში ორჯერ გვხვდება")
        seen.add(ext)
        label = " ".join(str(values.get(k) or "") for k in ("make", "model", "year")).strip()
        if errors:
            results.append(RowResult(n, ext, "error", errors, label))
            continue
        car = existing.get(ext)
        results.append(RowResult(n, ext, "updated" if car else "new", [], label))
        to_save.append((car, values))

    if commit:
        with transaction.atomic():
            for car, values in to_save:
                if car is None:
                    Car.objects.create(organization=organization, **values)
                else:
                    for k, v in values.items():
                        setattr(car, k, v)
                    car.save()

    return ImportResult(results, committed=commit)


def template_workbook():
    """The .xlsx template dealers fill in, with dropdowns for fixed values."""
    wb = Workbook()
    ws = wb.active
    ws.title = "მანქანები"
    headers = list(COLUMNS)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(h) + 4)

    for f, choices in (
        ("body", Car.Body.choices),
        ("fuel", Car.Fuel.choices),
        ("transmission", Car.Transmission.choices),
    ):
        col = ws.cell(1, headers.index(next(h for h, v in COLUMNS.items() if v == f)) + 1).column_letter
        dv = DataValidation(
            type="list", formula1='"' + ",".join(str(l) for _, l in choices) + '"', allow_blank=True
        )
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{settings.EXCEL_MAX_ROWS + 1}")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
